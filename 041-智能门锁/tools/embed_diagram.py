# -*- coding: utf-8 -*-
"""Embed task_queue_dataflow diagram into pptx (new slide) and docx (after 4.2)."""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

from docx import Document
from docx.shared import Inches as DInches, Pt as DPt, RGBColor as DRGB
from docx.enum.text import WD_ALIGN_PARAGRAPH

PNG = r"C:/Users/71563/Desktop/block/docs/task_queue_dataflow.png"
PPTX = r"C:/Users/71563/Desktop/block/docs/code_walkthrough.pptx"
DOCX = r"C:/Users/71563/Desktop/block/docs/智能门锁项目文档案例.docx"

def set_run_font(run, size, bold, color):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = "Microsoft YaHei"
    rPr = run._r.get_or_add_rPr()
    latin = rPr.find(qn('a:latin'))
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        if latin is not None:
            latin.addnext(ea)
        else:
            rPr.append(ea)
    ea.set('typeface', 'Microsoft YaHei')
    cs = rPr.find(qn('a:cs'))
    if cs is None:
        cs = rPr.makeelement(qn('a:cs'), {})
        ea.addnext(cs)
    cs.set('typeface', 'Microsoft YaHei')

# ---------------- PPTX ----------------
prs = Presentation(PPTX)
SW, SH = prs.slide_width, prs.slide_height
blank = prs.slide_masters[0].slide_layouts[6]
slide = prs.slides.add_slide(blank)

# background full-bleed
bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
bg.fill.solid(); bg.fill.fore_color.rgb = RGBColor.from_string("F5F7FA")
bg.line.fill.background()
bg.shadow.inherit = False

# header band
band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Inches(1.15))
band.fill.solid(); band.fill.fore_color.rgb = RGBColor.from_string("1F6FB2")
band.line.fill.background()
band.shadow.inherit = False

# title textbox (kicker + title)
tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.12), Inches(12.33), Inches(0.95))
tf = tb.text_frame
tf.word_wrap = True
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]
r1 = p.add_run(); r1.text = "总览  "
set_run_font(r1, 12, True, "CFE6FF")
r2 = p.add_run(); r2.text = "任务 — 队列 — 数据流（5 任务 + 6 队列）"
set_run_font(r2, 26, True, "FFFFFF")

# image, centered, below header
img_h = Inches(5.72)
pic = slide.shapes.add_picture(PNG, 0, Inches(1.34), height=img_h)
pic.left = int((SW - pic.width) / 2)
print("pic:", Emu(pic.left).inches, Emu(pic.top).inches, Emu(pic.width).inches, Emu(pic.height).inches)

# caption
cap = slide.shapes.add_textbox(Inches(0.5), Inches(7.12), Inches(12.33), Inches(0.32))
cf = cap.text_frame; cf.word_wrap = True
cp = cf.paragraphs[0]; cp.alignment = PP_ALIGN.CENTER
cr = cp.add_run()
cr.text = "access_task 被 4 条队列指向 = 唯一决策中心；所有跨任务通信均经 xQueueCreateStatic 静态队列，无动态分配"
set_run_font(cr, 11, False, "5A6B7B")

prs.save(PPTX)
# verify
prs2 = Presentation(PPTX)
print("pptx slides now:", len(prs2.slides._sldIdLst))
last = prs2.slides[-1]
print("last slide shapes:", [(sh.shape_type, sh.name) for sh in last.shapes])

# ---------------- DOCX ----------------
doc = Document(DOCX)
target = None
for par in doc.paragraphs:
    if par.text.strip().startswith("4.3"):
        target = par
        break
assert target is not None, "4.3 heading not found"

img_p = target.insert_paragraph_before()
img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = img_p.add_run()
run.add_picture(PNG, width=DInches(5.9))

cap_p = target.insert_paragraph_before()
cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
crun = cap_p.add_run("图 4-1  任务 — 队列 — 数据流总览（5 任务 + 6 队列，access_task 为唯一决策中心）")
crun.font.size = DPt(9)
crun.font.name = "Microsoft YaHei"
crun.font.color.rgb = DRGB(0x5A, 0x6B, 0x7B)
rpr = crun._element.get_or_add_rPr()
rf = rpr.find(qn('w:rFonts'))
if rf is None:
    rf = rpr.makeelement(qn('w:rFonts'), {})
    rpr.insert(0, rf)
rf.set(qn('w:eastAsia'), 'Microsoft YaHei')

doc.save(DOCX)
# verify
doc2 = Document(DOCX)
import zipfile
with zipfile.ZipFile(DOCX) as z:
    media = [n for n in z.namelist() if n.startswith("word/media/")]
print("docx media files:", media)
print("done")
